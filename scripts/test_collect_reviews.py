import unittest,sys,json,tempfile
from pathlib import Path
sys.path.insert(0,str(Path(__file__).parent))
from collect_reviews import parse_item_url, parse_response, normalize_comments, merge_rows, CollectionError, read_checkpoint, write_checkpoint, projected_rows
class CollectorTests(unittest.TestCase):
 def test_url_html_entity_and_source_selection(self):
  r=parse_item_url('https://item.taobao.com/item.htm?id=841836517132&amp;skuId=5');self.assertEqual(r['id'],'841836517132');self.assertIn('&skuId=5',r['url'])
  for u in ['javascript:alert(1)','https://evil.test/?id=1','https://item.taobao.com/item.htm?id=1&id=2']:
   with self.assertRaises(CollectionError):parse_item_url(u)
 def test_json_wrappers_without_eval_and_challenge(self):
  for t in ['({"comments":[],"total":0})','jsonp123({"comments":[],"total":0});','{"comments":[],"total":0}']:
   self.assertEqual(parse_response(t)['comments'],[])
  for t in ['callback({});alert(1)','{"url":"https://login.taobao.com/"}','<html>verify</html>','{"ret":["SUCCESS::ok","FAIL_SYS_SESSION_EXPIRED::bad"],"comments":[]}','{"ret":"bad","comments":[]}']:
   with self.assertRaises(CollectionError):parse_response(t)
 def test_media_initial_append_source_and_exact_numeric_id(self):
  d=[{'rateId':1311234567890123456,'content':'A &amp; B','date':'date','auction':{'sku':'C:D;E'},'photos':[{'url':'//img.alicdn.com/a.jpg_400x400.jpg?x=1'}],'video':{'videoId':3,'cloudVideoUrl':''},'append':{'appendId':2,'content':'追评','photos':[]},'appendList':[{'appendId':2,'content':'追评','photos':[]}]}]
  r=normalize_comments(d,{'id':'1'},1,'https://rate.taobao.com/');self.assertEqual(len(r),2);self.assertEqual(r[0]['id'],'1:1311234567890123456');self.assertEqual(r[0]['text'],'A & B');self.assertEqual(r[0]['images'][0]['url'],'https://img.alicdn.com/a.jpg?x=1');self.assertEqual(r[0]['images'][0]['sourceUrl'],'https://img.alicdn.com/a.jpg_400x400.jpg?x=1');self.assertIsNone(r[1]['date']);self.assertEqual(r[1]['parentId'],r[0]['id']);self.assertIsNone(r[0]['videos'][0]['url']);self.assertEqual(merge_rows(r,r),r)
 def test_invalid_identity_and_response_not_empty(self):
  for identity in [{},[],True,1.3,None]:
   with self.assertRaises(CollectionError):normalize_comments([{'rateId':identity,'photos':[]}],{'id':'1'},1,'u')
  with self.assertRaises(CollectionError):parse_response('{"data":{}}')
 def test_checkpoint_scoped_and_digest_verified(self):
  with tempfile.TemporaryDirectory() as d:
   p=Path(d)/'one.json';item={'id':'1','sellerId':'2'};text='({"comments":[],"total":0})';write_checkpoint(p,item,1,'image','u',text)
   self.assertEqual(read_checkpoint(p,item,1,'image')['rawText'],text)
   with self.assertRaises(CollectionError):read_checkpoint(p,{'id':'9','sellerId':'2'},1,'image')
   v=json.loads(p.read_text());v['rawText']='changed';p.write_text(json.dumps(v));
   with self.assertRaises(CollectionError):read_checkpoint(p,item,1,'image')
 def test_projection_does_not_drop_source_scope(self):
  rows=[{'id':'1','videos':[]},{'id':'2','videos':[{'url':None,'sourceId':'3'}]}];self.assertEqual(projected_rows(rows,'video'),[rows[1]]);self.assertEqual(projected_rows(rows,'all'),rows)
if __name__=='__main__':unittest.main()

class IdentityEndpointTests(unittest.TestCase):
 def test_missing_public_seller_keeps_id_and_omits_user_parameter(self):
  from collect_reviews import rate_url, validate_page_identity
  u=rate_url({'id':'123','sellerId':None},1,'image');self.assertNotIn('userNumId=',u)
  validate_page_identity([{'auction':{'aucNumId':'123'}}],{'id':'123'})
  with self.assertRaises(CollectionError):validate_page_identity([{'auction':{'aucNumId':'456'}}],{'id':'123'})

class RealFormatBoundaryTests(unittest.TestCase):
 def test_heic_original_source_boundary_preserves_queries(self):
  from collect_reviews import original_image_url
  self.assertEqual(original_image_url('https://img.alicdn.com/rate.heic_400x400.jpg?x=1'),'https://img.alicdn.com/rate.heic?x=1')
  self.assertEqual(original_image_url('https://evil.test/rate.heic_400x400.jpg'),'https://evil.test/rate.heic_400x400.jpg')

class RecoveryTests(unittest.TestCase):
 def test_online_resume_refetches_explicit_failed_page_preserving_original(self):
  from collect_reviews import collect_pages
  with tempfile.TemporaryDirectory() as d:
   raw=Path(d)/'raw';raw.mkdir();item={'id':'1','sellerId':None};cp=raw/'page-0001.json';failed='{"url":"https://login.taobao.com/"}';write_checkpoint(cp,item,1,'image','u',failed);calls=[]
   def fetch(url):calls.append(url);return {'status':200,'rawText':'({"comments":[{"rateId":1,"photos":[],"auction":{"aucNumId":"1"}}],"total":1})'}
   r=collect_pages(item,'image',raw,fetch,delay=0);self.assertTrue(r['coverage']['complete']);self.assertEqual(len(calls),1);self.assertTrue(any(failed==json.loads(x.read_text())['rawText'] for x in (raw/'failed-attempts').glob('*.json')))
 def test_offline_failed_page_does_not_fetch_or_hide_rejection(self):
  from collect_reviews import collect_pages
  with tempfile.TemporaryDirectory() as d:
   raw=Path(d)/'raw';raw.mkdir();item={'id':'1','sellerId':None};write_checkpoint(raw/'page-0001.json',item,1,'image','u','{"url":"https://login.taobao.com/"}');calls=[]
   with self.assertRaises(CollectionError):collect_pages(item,'image',raw,lambda u:calls.append(u),delay=0,online=False)
   self.assertEqual(calls,[])


class NoVideoCliTests(unittest.TestCase):
 def test_offline_dispatch_excludes_export_videos_without_browser_or_source_loss(self):
  import subprocess
  from openpyxl import load_workbook
  from collect_reviews import rate_url
  with tempfile.TemporaryDirectory() as folder:
   out=Path(folder);item={'id':'123','sellerId':None,'url':'https://item.taobao.com/item.htm?id=123'}
   (out/'session.json').write_text(json.dumps({'item':item,'filter':'image'}),encoding='utf-8')
   raw=json.dumps({'comments':[{'rateId':1234567890123456789,'content':'Synthetic initial','auction':{'aucNumId':'123','sku':'Synthetic SKU'},'photos':[{'url':'https://example.invalid/picture.jpg'}],'video':{'cloudVideoUrl':'https://example.invalid/initial.mp4'},'append':{'appendId':9876543210987654321,'content':'Synthetic append','video':{'videoId':7}}}],'total':1})
   checkpoint=out/'raw/page-0001.json';write_checkpoint(checkpoint,item,1,'image',rate_url(item,1,'image'),raw)
   original=checkpoint.read_bytes()
   script=Path(__file__).with_name('collect.py')
   done=subprocess.run([sys.executable,'-X','utf8','-B',str(script),'--platform','taobao','--from-raw','--output-dir',str(out),'--exclude-videos','--filename','no-video.xlsx'],capture_output=True,text=True,encoding='utf-8')
   self.assertEqual(done.returncode,0,done.stderr)
   self.assertEqual(json.loads(done.stdout.splitlines()[-1])['video_policy'],'excluded')
   result=json.loads((out/'reviews.json').read_text(encoding='utf-8'))
   receipt=json.loads((out/'receipt.json').read_text(encoding='utf-8'))
   self.assertEqual(checkpoint.read_bytes(),original)
   self.assertEqual(len(result['rows']),2)
   self.assertEqual(result['rows'][0]['id'],'123:1234567890123456789')
   self.assertEqual(result['rows'][1]['parentId'],result['rows'][0]['id'])
   self.assertEqual(len(result['rows'][0]['videos']),1)
   self.assertIsNone(result['rows'][1]['videos'][0]['url'])
   self.assertTrue(result['coverage']['complete'])
   self.assertEqual(receipt['video_policy'],'excluded')
   self.assertEqual(receipt['excluded_video_url_count'],1)
   workbook=load_workbook(out/'no-video.xlsx')
   try:
    self.assertEqual(workbook.sheetnames,['全部'])
    self.assertEqual(workbook.active.max_row,3)
    self.assertEqual(workbook.active.max_column,5)
    self.assertTrue(all(workbook.active.cell(row,3).value is None for row in (2,3)))
   finally:workbook.close()
