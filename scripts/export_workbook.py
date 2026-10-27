#!/usr/bin/env python3
"""Export normalized Taobao reviews without fetching or interpreting their content.

Requires Python 3.12 and preinstalled openpyxl (tested with 3.1.5); installs nothing.
API: export_workbook(input_path, output_path, template_path=None) -> dict.
Parser: parse_normalized_input(input_path) -> NormalizedInput (validated original JSON).

Required JSON shape (extra metadata is allowed):
  item: {id: str | int, url: str}
  rows: [{id: str | int | null, parentId: str | int | null,
          type: "initial" | "append", text: str | null, sku: str | null,
          date: str | null, images: [{url: str}], videos: [{url: str | null}]}]
  coverage: object (opaque caller metadata; no completeness is inferred)
All row fields are required. Order, whitespace and URLs are retained; null dates
are blank, not inferred. Missing video URLs produce no invented URL and are counted
in <output-stem>.coverage.json. Empty strings and null display as blank cells except
SKU: null/"" display as "未显示sku", a presentation placeholder, not an observed SKU.
The normalized JSON/domain retains its original value; explicit SKU labels and
whitespace-only strings are unchanged.
Templates supply structure/styles only: one sheet named 全部 and the five headers.
No reference data, comments, links, embedded media or workbook metadata are copied.
"""
from __future__ import annotations

import argparse
from copy import copy
import json
from io import BytesIO
from pathlib import Path
import os, subprocess, shutil, tempfile
import sys
from typing import Literal, TypedDict
from xml.etree import ElementTree as ET
from zipfile import BadZipFile, ZIP_DEFLATED, ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.worksheet._writer import WorksheetWriter
from openpyxl.writer.excel import ExcelWriter

HEADERS = ("内容", "图片", "视频", "SKU", "评论时间")
SHEET_NAME = "全部"
MISSING_SKU_LABEL = "未显示sku"
DEFAULT_TEMPLATE = Path(__file__).resolve().parent.parent / "assets" / "workbook-template.xlsx"
XML_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
STYLE_FIELDS = ("font", "fill", "border", "alignment", "number_format", "protection")


class Item(TypedDict):
    id: str | int
    url: str


class Image(TypedDict):
    url: str


class Video(TypedDict):
    url: str | None


class ReviewRow(TypedDict):
    id: str | int | None
    parentId: str | int | None
    type: Literal["initial", "append"]
    text: str | None
    sku: str | None
    date: str | None
    images: list[Image]
    videos: list[Video]


class NormalizedInput(TypedDict):
    item: Item
    rows: list[ReviewRow]
    coverage: dict


def _require(condition: bool, field: str, expectation: str) -> None:
    if not condition:
        raise ValueError(f"{field}: expected {expectation}")


def _object(value, field: str, required: tuple[str, ...]) -> None:
    _require(isinstance(value, dict), field, "object")
    for key in required:
        _require(key in value, f"{field}.{key}", "required field")


def _identifier(value, field: str, nullable: bool = False) -> None:
    _require(type(value) in (str, int) or (nullable and value is None), field,
             "string or integer" + (" or null" if nullable else ""))


def _reject_constant(value: str):
    raise ValueError(f"Non-standard JSON constant: {value}")


def parse_normalized_input(input_path) -> NormalizedInput:
    """Validate at the owning JSON boundary; return source values without normalization."""
    data = json.loads(Path(input_path).read_text(encoding="utf-8-sig"), parse_constant=_reject_constant)
    _object(data, "$", ("item", "rows", "coverage"))
    _object(data["item"], "item", ("id", "url"))
    _identifier(data["item"]["id"], "item.id")
    _require(isinstance(data["item"]["url"], str), "item.url", "string")
    _require(isinstance(data["coverage"], dict), "coverage", "object")
    _require(isinstance(data["rows"], list), "rows", "array")
    _require(len(data["rows"]) <= 1048575, "rows", "at most 1,048,575 data rows (Excel limit)")
    for i, row in enumerate(data["rows"]):
        field = f"rows[{i}]"
        _object(row, field, ("id", "parentId", "type", "text", "sku", "date", "images", "videos"))
        for key in ("id", "parentId"):
            _identifier(row[key], f"{field}.{key}", nullable=True)
        _require(row["type"] in ("initial", "append"), f"{field}.type", "initial or append")
        for key in ("text", "sku", "date"):
            _require(row[key] is None or isinstance(row[key], str), f"{field}.{key}", "string or null")
        for key in ("images", "videos"):
            _require(isinstance(row[key], list), f"{field}.{key}", "array")
            for j, media in enumerate(row[key]):
                path = f"{field}.{key}[{j}]"
                _object(media, path, ("url",))
                _require(isinstance(media["url"], str) or (key == "videos" and media["url"] is None),
                         f"{path}.url", "string" + (" or null" if key == "videos" else ""))
    return data


def _excel_string(value: str | None, field: str) -> None:
    if value is None:
        return
    # openpyxl would silently truncate long strings; fail before any output write.
    _require(len(value) <= 32767, field, "at most 32,767 characters (Excel cell limit)")
    for char in value:
        code = ord(char)
        _require(code in (9, 10, 13) or 32 <= code <= 0xD7FF or 0xE000 <= code <= 0xFFFD
                 or 0x10000 <= code <= 0x10FFFF, field, "XML 1.0 representable text")


def _project_rows(data: NormalizedInput, *, exclude_videos: bool = False) -> list[tuple]:
    """Own workbook display rules without modifying normalized source facts."""
    _require(type(exclude_videos) is bool, "exclude_videos", "explicit boolean")
    projected = []
    for i, row in enumerate(data["rows"]):
        values = (row["text"], ",".join(m["url"] for m in row["images"]),
                  "" if exclude_videos else ",".join(m["url"] for m in row["videos"] if m["url"] is not None),
                  MISSING_SKU_LABEL if row["sku"] in (None, "") else row["sku"], row["date"])
        for header, value in zip(HEADERS, values):
            _excel_string(value, f"rows[{i}].{header}")
        projected.append(values)
    return projected


def _styles(cell) -> dict:
    return {key: copy(getattr(cell, key)) for key in STYLE_FIELDS}


def _apply_style(target, style: dict) -> None:
    for key, value in style.items():
        setattr(target, key, copy(value))


def _blank_from_reference(reference_path):
    """Rebuild rather than deleting rows: residual strings/comments must not survive."""
    source = load_workbook(reference_path, keep_links=False)
    try:
        _require(source.sheetnames == [SHEET_NAME], "template", "only one sheet named 全部")
        sheet = source.active
        _require(sheet.max_column == 5 and tuple(c.value for c in sheet[1]) == HEADERS,
                 "template", "exact headers 内容/图片/视频/SKU/评论时间")
        _require(not sheet.merged_cells, "template", "unmerged five-column structure")
        target = Workbook()
        out = target.active
        out.title = SHEET_NAME
        # A fresh workbook also removes source metadata, embedded media and defined names.
        target._fonts[0] = copy(source._fonts[0])
        target._named_styles[0].font = copy(source._fonts[0])
        out.sheet_format = copy(sheet.sheet_format)
        for key, dimension in sheet.column_dimensions.items():
            new = copy(dimension)
            new.parent = out
            out.column_dimensions[key] = new
            _apply_style(new, _styles(dimension))
        if 1 in sheet.row_dimensions:
            dimension = sheet.row_dimensions[1]
            new = copy(dimension)
            new.parent = out
            out.row_dimensions[1] = new
            _apply_style(new, _styles(dimension))
        for col, header in enumerate(HEADERS, 1):
            cell = out.cell(1, col, header)
            cell.data_type = "s"
            _apply_style(cell, _styles(sheet.cell(1, col)))
        body_styles = [_styles(sheet.cell(2, col)) for col in range(1, 6)]
        return target, body_styles
    finally:
        source.close()


class _StringWorkbookWriter(ExcelWriter):
    """Serialization boundary for our fresh, single-sheet, string-only workbook."""

    def write_worksheet(self, worksheet):
        # A binary stream avoids et_xmlfile's Windows text-file LF -> CRLF
        # translation. Escape literal CR before XML readers normalize it to LF.
        # The fresh workbook has no drawings/tables/comments to serialize.
        with BytesIO() as buffer:
            writer = WorksheetWriter(worksheet, out=buffer)
            writer.write()
            worksheet._rels = writer._rels
            content = buffer.getvalue().replace(b"\r", b"&#13;")
            self._archive.writestr(worksheet.path[1:], content)
            self.manifest.append(worksheet)


def _save_workbook(workbook, output_path: Path) -> None:
    """Use the bundled sheet author when available; keep fallback for standalone hosts."""
    base=Path.home()/'.cache/codex-runtimes/codex-primary-runtime/dependencies/node'
    entry=base/'node_modules/@oai/artifact-tool/dist/artifact_tool.mjs'
    node=base/'bin/node.exe'
    if not node.is_file():node=Path(shutil.which('node') or '')
    preserve_cr=any(isinstance(c.value,str) and "\r" in c.value for row in workbook.active for c in row)
    # Current artifact exporter normalizes CRLF to LF; required literal source CR
    # preservation is a missing capability, use format-aware string writer for it.
    if entry.is_file() and node.is_file() and not preserve_cr:
        sheet=workbook.active
        font=lambda f:{'name':f.name or 'Calibri','size':f.sz or 12,'bold':bool(f.b),'italic':bool(f.i)}
        spec={'sheet':sheet.title,'values':[list(row) for row in sheet.values],'widths':[sheet.column_dimensions[c].width for c in 'ABCDE'],'defaultFont':font(workbook._fonts[0]),'headerFonts':[font(sheet.cell(1,c).font) for c in range(1,6)],'bodyFonts':[font(sheet.cell(2,c).font) for c in range(1,6)],'bodyWrap':[bool(sheet.cell(2,c).alignment.wrap_text) for c in range(1,6)],'headerHeight':sheet.row_dimensions[1].height}
        with tempfile.TemporaryDirectory(prefix='review-xlsx-') as d:
            source=Path(d)/'spec.json';source.write_text(json.dumps(spec,ensure_ascii=False,allow_nan=False),encoding='utf-8')
            result=subprocess.run([str(node),str(Path(__file__).with_name('artifact_workbook.mjs')),str(entry),str(source),str(output_path)],capture_output=True,text=True,encoding='utf-8',timeout=60)
            if result.returncode:raise ValueError('artifact workbook author failed: '+result.stderr[-1000:])
    else:
        with ZipFile(output_path, "w", ZIP_DEFLATED, allowZip64=True) as archive:
            _StringWorkbookWriter(workbook, archive).save()


def _verify_output(output_path: Path, values: list[tuple]) -> dict:
    workbook=load_workbook(output_path,data_only=False,keep_links=False)
    try:
        _require(workbook.sheetnames==[SHEET_NAME],'output','exactly expected sheet')
        sheet=workbook.active
        _require(sheet.max_column==5 and sheet.max_row==len(values)+1,'output','exact five-column dimensions and row count')
        for index,(actual,wanted) in enumerate(zip(sheet.values,[HEADERS,*values]),1):
            wanted=tuple(None if value=='' else value for value in wanted)
            _require(actual==wanted,f'output.row[{index}]','lossless readback')
        _require(all(cell.data_type!='f' for row in sheet for cell in row),'output','string data, not formulas')
    finally:workbook.close()
    with ZipFile(output_path) as archive:
        for name in archive.namelist():
            if name.startswith('xl/worksheets/') and name.endswith('.xml'):
                root=ET.fromstring(archive.read(name));_require(root.find(f'.//{{{XML_NS}}}f') is None,'output XML','no formula elements')
    return {'readback':True,'formula_free':True}

def export_workbook(input_path,output_path,template_path=None, *, exclude_videos: bool = False) -> dict:
    input_path=Path(input_path).resolve();output_path=Path(output_path).resolve();template_path=Path(template_path or DEFAULT_TEMPLATE).resolve()
    sidecar_path=output_path.with_suffix('.coverage.json')
    _require(output_path.suffix.lower()=='.xlsx','output','.xlsx extension')
    _require(output_path not in {input_path,template_path} and sidecar_path not in {input_path,template_path},'output','not input or template')
    data=parse_normalized_input(input_path);values=_project_rows(data,exclude_videos=exclude_videos);workbook,styles=_blank_from_reference(template_path)
    try:
        for index,row in enumerate(values,2):
            for col,value in enumerate(row,1):
                cell=workbook.active.cell(index,col,value);cell.data_type='s';_apply_style(cell,styles[col-1])
        _save_workbook(workbook,output_path)
    finally:workbook.close()
    validation=_verify_output(output_path,values)
    summary={'item':data['item'],'coverage':data['coverage'],'row_count':len(values),'initial_count':sum(row['type']=='initial' for row in data['rows']),'append_count':sum(row['type']=='append' for row in data['rows']),'missing_video_url_count':sum(m['url'] is None for row in data['rows'] for m in row['videos'])}
    video_policy={'video_policy':'excluded' if exclude_videos else 'included','excluded_video_url_count':sum(m['url'] is not None for row in data['rows'] for m in row['videos']) if exclude_videos else 0}
    summary.update(video_policy)
    sidecar_path.write_text(json.dumps(summary,ensure_ascii=False,indent=2,allow_nan=False)+'\n',encoding='utf-8')
    return {'output_path':str(output_path),'coverage_path':str(sidecar_path),'row_count':len(values),'sheet_name':SHEET_NAME,'columns':list(HEADERS),'validation':validation,**video_policy}

def main(argv=None) -> int:
    parser=argparse.ArgumentParser(description=__doc__);parser.add_argument('--input',required=True,type=Path);parser.add_argument('--output',required=True,type=Path);parser.add_argument('--exclude-videos',action='store_true',help='Leave video cells blank; preserve source rows and evidence');templates=parser.add_mutually_exclusive_group();templates.add_argument('--template',type=Path);templates.add_argument('--reference-file',type=Path);args=parser.parse_args(argv)
    try:result=export_workbook(args.input,args.output,args.reference_file or args.template,exclude_videos=args.exclude_videos)
    except (ValueError,OSError,KeyError,BadZipFile,InvalidFileException,ET.ParseError) as error:
        print(f'export error: {error}',file=sys.stderr);return 1
    print(json.dumps(result,ensure_ascii=True));return 0
if __name__=='__main__':raise SystemExit(main())
