"""Synthetic fixtures only. Run with Python 3.12 -B using this absolute path."""
import copy
import importlib.util
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment

SCRIPT = Path(__file__).with_name("export_workbook.py")
TEMPLATE = SCRIPT.parent.parent / "assets" / "workbook-template.xlsx"
HEADERS = ["内容", "图片", "视频", "SKU", "评论时间"]
NS = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}


def synthetic_payload():
    return {
        "item": {"id": "synthetic-item", "url": "https://example.invalid/item?id=sample"},
        "rows": [
            {"id": "r1", "parentId": None, "type": "initial",
             "text": "<b>造数据 &amp; 不解析</b>\n 第二行 ", "sku": " 蓝色 / M ",
             "date": "2026-01-02 03:04:05",
             "images": [{"url": "//example.invalid/a?x=1&y=2"}, {"url": "https://example.invalid/b,c"}],
             "videos": [{"url": None}, {"url": "https://example.invalid/v?token=synthetic"}]},
            {"id": "r2", "parentId": "r1", "type": "append", "text": "=1+1",
             "sku": None, "date": None, "images": [], "videos": [{"url": None}]},
            {"id": "r3", "parentId": None, "type": "initial", "text": "+SUM(1,2)",
             "sku": "@sku", "date": " ", "images": [], "videos": []},
            {"id": "r4", "parentId": None, "type": "initial", "text": "-1+2",
             "sku": "", "date": "", "images": [], "videos": []},
        ],
        "coverage": {"complete": False, "reason": "synthetic partial coverage", "pages": 2,
                     "unknownVideoUrls": 2, "nested": {"callerChoice": "unchanged"}},
    }


class ExportWorkbookTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        spec = importlib.util.spec_from_file_location("export_workbook", SCRIPT)
        cls.api = importlib.util.module_from_spec(spec)
        sys.modules[spec.name] = cls.api
        spec.loader.exec_module(cls.api)

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.input = self.root / "input.json"
        self.output = self.root / "output.xlsx"
        self.payload = synthetic_payload()
        self.write_input()

    def write_input(self):
        self.input.write_text(json.dumps(self.payload, ensure_ascii=False), encoding="utf-8")

    def export(self, template=None):
        self.write_input()
        return self.api.export_workbook(self.input, self.output, template_path=template)

    def test_exclude_videos_is_display_only_and_preserves_source_scope(self):
        self.payload["sourceRows"] = copy.deepcopy(self.payload["rows"])
        self.write_input()
        original = self.input.read_bytes()
        parsed = self.api.parse_normalized_input(self.input)
        source_facts = copy.deepcopy(parsed)
        projected = self.api._project_rows(parsed, exclude_videos=True)
        self.assertEqual(parsed, source_facts)
        self.assertEqual(len(projected), len(parsed["rows"]))
        self.assertTrue(all(row[2] == "" for row in projected))
        receipt = self.api.export_workbook(self.input, self.output, exclude_videos=True)
        self.assertEqual(self.input.read_bytes(), original)
        workbook = load_workbook(self.output)
        self.addCleanup(workbook.close)
        self.assertEqual(workbook.sheetnames, ["全部"])
        self.assertEqual(workbook.active.max_column, 5)
        self.assertEqual(list(workbook.active.values),
                         [tuple(HEADERS), *[tuple(None if value == "" else value for value in row)
                                           for row in projected]])
        summary = json.loads(Path(receipt["coverage_path"]).read_text(encoding="utf-8"))
        self.assertEqual(summary["coverage"], self.payload["coverage"])
        self.assertEqual(summary["video_policy"], "excluded")
        self.assertEqual(summary["excluded_video_url_count"], 1)
        self.assertEqual(receipt["video_policy"], "excluded")
        self.assertEqual(receipt["excluded_video_url_count"], 1)
        self.assertEqual(summary["missing_video_url_count"], 2)
        self.assertEqual(receipt["row_count"], len(self.payload["rows"]))

    def test_exclude_videos_requires_explicit_boolean(self):
        for value in ("false", None, 0, 1):
            with self.subTest(value=value), self.assertRaises(ValueError):
                self.api.export_workbook(self.input, self.output, exclude_videos=value)
        self.assertFalse(self.output.exists())

    def test_cli_exclude_videos_keeps_video_column_and_all_rows(self):
        original = self.input.read_bytes()
        done = subprocess.run([sys.executable, "-B", str(SCRIPT), "--input", str(self.input),
                               "--output", str(self.output), "--exclude-videos"],
                              capture_output=True, text=True, encoding="utf-8")
        self.assertEqual(done.returncode, 0, done.stderr)
        receipt = json.loads(done.stdout)
        self.assertEqual(receipt["video_policy"], "excluded")
        self.assertEqual(self.input.read_bytes(), original)
        workbook = load_workbook(self.output)
        self.addCleanup(workbook.close)
        self.assertEqual(workbook.active.max_row, len(self.payload["rows"]) + 1)
        self.assertEqual([cell.value for cell in workbook.active[1]], HEADERS)
        self.assertTrue(all(workbook.active.cell(row, 3).value is None
                            for row in range(2, workbook.active.max_row + 1)))

    def test_packaged_template_is_header_only_with_reference_style(self):
        with ZipFile(TEMPLATE) as z:
            self.assertFalse(any(n.startswith("xl/media/") or n == "xl/sharedStrings.xml" for n in z.namelist()))
            xml = ET.fromstring(z.read("xl/worksheets/sheet1.xml"))
            self.assertFalse(xml.findall(".//s:f", NS))
            self.assertEqual([c.attrib["r"] for c in xml.findall(".//s:c", NS)], ["A1", "B1", "C1", "D1", "E1"])
        w = load_workbook(TEMPLATE)
        self.addCleanup(w.close)
        self.assertEqual(w.sheetnames, ["全部"])
        self.assertEqual(list(w.active.values), [tuple(HEADERS)])
        for c in w.active[1]:
            self.assertEqual((c.font.name, c.font.sz), ("Calibri", 12))
        self.assertEqual((w._fonts[0].name, w._fonts[0].sz), ("Calibri", 12))
        for col in "ABCDE":
            self.assertEqual(w.active.column_dimensions[col].width, 36.6640625)

    def test_missing_sku_is_display_placeholder_without_changing_source_facts(self):
        skus = [None, "", " ", "未显示sku", " 未显示sku ", "=explicit raw label", "造数据原始SKU"]
        self.payload["rows"] = []
        for index, sku in enumerate(skus):
            row = copy.deepcopy(synthetic_payload()["rows"][0])
            row.update(id=f"synthetic-sku-{index}", sku=sku)
            self.payload["rows"].append(row)
        self.write_input()
        original_bytes = self.input.read_bytes()
        parsed = self.api.parse_normalized_input(self.input)
        original_facts = copy.deepcopy(parsed)
        self.api._project_rows(parsed)
        self.assertEqual(parsed, original_facts)
        self.api.export_workbook(self.input, self.output)
        self.assertEqual(self.input.read_bytes(), original_bytes)
        self.assertEqual(self.api.parse_normalized_input(self.input), original_facts)
        w = load_workbook(self.output)
        self.addCleanup(w.close)
        self.assertEqual([w.active.cell(i + 2, 4).value for i in range(len(skus))],
                         ["未显示sku", "未显示sku", *skus[2:]])
        self.assertEqual(w.sheetnames, ["全部"])
        self.assertEqual(w.active.max_column, 5)
        for index in range(len(skus)):
            cell = w.active.cell(index + 2, 4)
            self.assertEqual(cell.data_type, "s")
            self.assertEqual((cell.font.name, cell.font.sz), ("Calibri", 12))
        self.assertEqual(w.active.column_dimensions["D"].width, 36.6640625)
        sidecar = json.loads(self.output.with_suffix(".coverage.json").read_text(encoding="utf-8"))
        self.assertEqual(sidecar["coverage"], self.payload["coverage"])

    def test_preserves_five_columns_raw_text_urls_and_append_order(self):
        result = self.export()
        w = load_workbook(self.output)
        self.addCleanup(w.close)
        self.assertEqual(w.sheetnames, ["全部"])
        self.assertEqual(list(w.active.values), [tuple(HEADERS),
            (self.payload["rows"][0]["text"], "//example.invalid/a?x=1&y=2,https://example.invalid/b,c",
             "https://example.invalid/v?token=synthetic", " 蓝色 / M ", "2026-01-02 03:04:05"),
            ("=1+1", None, None, "未显示sku", None),
            ("+SUM(1,2)", None, None, "@sku", " "),
            ("-1+2", None, None, "未显示sku", None)])
        self.assertEqual(result["row_count"], 4)
        self.assertEqual(result["validation"], {"readback": True, "formula_free": True})
        for row in w.active:
            for c in row:
                self.assertNotEqual(c.data_type, "f")
                self.assertEqual((c.font.name, c.font.sz), ("Calibri", 12))

    def test_xml_contains_strings_not_formulas_or_hyperlinks(self):
        self.export()
        with ZipFile(self.output) as z:
            for name in z.namelist():
                if name.startswith("xl/worksheets/") and name.endswith(".xml"):
                    root = ET.fromstring(z.read(name))
                    self.assertFalse(root.findall(".//s:f", NS))
                    self.assertFalse(root.findall(".//s:hyperlink", NS))
            root = ET.fromstring(z.read("xl/worksheets/sheet1.xml"))
            cell = root.find(".//s:c[@r='A3']", NS)
            self.assertIn(cell.attrib["t"], ("inlineStr", "s", "str"))
            literal=cell.find(".//s:t", NS)
            if literal is None:literal=cell.find("s:v", NS)
            self.assertEqual(literal.text, "=1+1")

    def test_coverage_sidecar_preserves_caller_metadata_without_extra_sheets(self):
        result = self.export()
        sidecar = self.output.with_suffix(".coverage.json")
        data = json.loads(sidecar.read_text(encoding="utf-8"))
        self.assertEqual(data["coverage"], self.payload["coverage"])
        self.assertEqual(data["item"], self.payload["item"])
        self.assertEqual(data["row_count"], 4)
        self.assertEqual(data["missing_video_url_count"], 2)
        self.assertEqual(Path(result["coverage_path"]), sidecar)

    def test_empty_rows_still_exports_header_only(self):
        self.payload["rows"] = []
        self.export()
        w = load_workbook(self.output)
        self.addCleanup(w.close)
        self.assertEqual(list(w.active.values), [tuple(HEADERS)])

    def test_custom_reference_copies_structure_but_not_data_or_comments(self):
        ref = self.root / "reference.xlsx"
        w = Workbook()
        s = w.active
        s.title = "全部"
        s.append(HEADERS)
        s.append(["SYNTHETIC_SENTINEL", None, None, "=2+2", "SYNTHETIC_DATE"])
        s.column_dimensions["A"].width = 52
        s.row_dimensions[1].height = 25
        s["A1"].font = Font(name="Arial", size=15, bold=True)
        s["A2"].alignment = Alignment(wrap_text=True)
        from openpyxl.comments import Comment
        s["A1"].comment = Comment("SYNTHETIC_SECRET", "synthetic")
        w.save(ref)
        before = ref.read_bytes()
        self.payload["rows"] = self.payload["rows"][:1]
        self.export(ref)
        self.assertEqual(ref.read_bytes(), before)
        w = load_workbook(self.output)
        self.addCleanup(w.close)
        self.assertEqual(w.active.column_dimensions["A"].width, 52)
        self.assertEqual(w.active.row_dimensions[1].height, 25)
        self.assertEqual(w.active["A1"].font.name, "Arial")
        self.assertTrue(w.active["A2"].alignment.wrap_text)
        self.assertIsNone(w.active["A1"].comment)
        with ZipFile(self.output) as z:
            for n in z.namelist():
                self.assertNotIn(b"SYNTHETIC_SENTINEL", z.read(n))
                self.assertNotIn(b"SYNTHETIC_SECRET", z.read(n))

    def test_rejects_incompatible_reference_without_writing_output(self):
        for extra_sheet in (False, True):
            with self.subTest(extra_sheet=extra_sheet):
                ref = self.root / "bad.xlsx"
                w = Workbook()
                w.active.title = "全部"
                w.active.append(HEADERS if extra_sheet else ["wrong"])
                if extra_sheet:
                    w.create_sheet("extra")
                w.save(ref)
                with self.assertRaises(ValueError):
                    self.export(ref)
                self.assertFalse(self.output.exists())

    def test_parser_rejects_invalid_schema_without_writing_output(self):
        original = copy.deepcopy(self.payload)
        mutations = [
            lambda p: p.update(rows={}),
            lambda p: p.update(coverage=[]),
            lambda p: p["item"].update(url=123),
            lambda p: p["rows"][0].update(type="unknown"),
            lambda p: p["rows"][0].update(text=42),
            lambda p: p["rows"][0].update(date=42),
            lambda p: p["rows"][0].update(images=[{"url": None}]),
            lambda p: p["rows"][0].update(videos=[{}]),
            lambda p: p["rows"][0].pop("parentId"),
        ]
        for mutate in mutations:
            self.payload = copy.deepcopy(original)
            mutate(self.payload)
            with self.subTest(payload=self.payload), self.assertRaises(ValueError):
                self.export()
            self.assertFalse(self.output.exists())

    def test_rejects_excel_unrepresentable_text_instead_of_truncating(self):
        for bad in ("x" * 32768, "x\x00y"):
            with self.subTest(length=len(bad)):
                self.payload["rows"][0]["text"] = bad
                with self.assertRaises(ValueError):
                    self.export()
                self.assertFalse(self.output.exists())

    def test_protects_input_and_template_from_output_collision(self):
        for target in (self.input, TEMPLATE):
            with self.subTest(target=target):
                before = target.read_bytes()
                with self.assertRaises(ValueError):
                    self.api.export_workbook(self.input, target)
                self.assertEqual(target.read_bytes(), before)

    def test_parser_entry_preserves_source_payload(self):
        self.assertEqual(self.api.parse_normalized_input(self.input), self.payload)

    def test_cli_default_template_and_reference_option(self):
        for option in ([], ["--reference-file", str(TEMPLATE)], ["--template", str(TEMPLATE)]):
            with self.subTest(option=option):
                done = subprocess.run([sys.executable, "-B", str(SCRIPT), "--input", str(self.input),
                                       "--output", str(self.output), *option], capture_output=True, text=True,
                                      encoding="utf-8", timeout=15)
                self.assertEqual(done.returncode, 0, done.stderr)
                self.assertEqual(json.loads(done.stdout)["row_count"], 4)

    def test_cli_invalid_template_returns_error_without_traceback(self):
        ref = self.root / "not-a-workbook.xlsx"
        ref.write_text("synthetic invalid workbook", encoding="utf-8")
        done = subprocess.run([sys.executable, "-B", str(SCRIPT), "--input", str(self.input),
                               "--output", str(self.output), "--template", str(ref)],
                              capture_output=True, text=True, encoding="utf-8", timeout=15)
        self.assertNotEqual(done.returncode, 0)
        self.assertNotIn("Traceback", done.stderr)
        self.assertFalse(self.output.exists())

    def test_blank_input_row_keeps_sku_placeholder_and_formula_like_media_as_string(self):
        self.payload["rows"] = [
            {"id": "blank", "parentId": None, "type": "initial", "text": None,
             "sku": None, "date": None, "images": [], "videos": []},
            {"id": "unicode", "parentId": None, "type": "initial", "text": "造数据😀\r\n尾行",
             "sku": "=sku", "date": "=date", "images": [{"url": "=image"}],
             "videos": [{"url": "=video"}]}]
        self.export()
        w = load_workbook(self.output)
        self.addCleanup(w.close)
        self.assertEqual(list(w.active.values)[1:], [
            (None, None, None, "未显示sku", None),
            ("造数据😀\r\n尾行", "=image", "=video", "=sku", "=date")])
        self.assertTrue(all(c.data_type != "f" for row in w.active for c in row))

    def test_iso_date_and_numeric_sku_remain_literal_strings(self):
        self.payload['rows']=self.payload['rows'][:1]
        self.payload['rows'][0]['date']='2026-10-01T00:00:00+08:00'
        self.payload['rows'][0]['sku']='00123'
        self.export()
        w=load_workbook(self.output)
        self.addCleanup(w.close)
        self.assertEqual(w.active['E2'].value,'2026-10-01T00:00:00+08:00')
        self.assertEqual(w.active['D2'].value,'00123')

    def test_cli_invalid_input_returns_error_without_traceback(self):
        self.input.write_text('{"rows":', encoding="utf-8")
        done = subprocess.run([sys.executable, "-B", str(SCRIPT), "--input", str(self.input),
                               "--output", str(self.output)], capture_output=True, text=True,
                              encoding="utf-8", timeout=15)
        self.assertNotEqual(done.returncode, 0)
        self.assertNotIn("Traceback", done.stderr)
        self.assertFalse(self.output.exists())


if __name__ == "__main__":
    unittest.main(verbosity=2)
